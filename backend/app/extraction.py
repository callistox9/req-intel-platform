import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


def extract_pdf(data: bytes) -> list[dict[str, Any]]:
    """Extract selectable text from each PDF page."""
    reader = PdfReader(io.BytesIO(data))

    return [
        {
            "page": page_number,
            "text": (page.extract_text() or "").strip(),
        }
        for page_number, page in enumerate(reader.pages, start=1)
    ]


def extract_docx(data: bytes) -> list[dict[str, Any]]:
    """Extract paragraphs and table rows from a DOCX file."""
    document = Document(io.BytesIO(data))
    lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]

            if any(cells):
                lines.append(
                    f"[Table {table_number}, row {row_number}] " + " | ".join(cells)
                )

    return [{"page": None, "text": "\n".join(lines)}]


def _merge_continuation_lines(lines: list[str]) -> list[str]:
    """Merge wrapped requirement lines before classifying them."""
    merged: list[str] = []

    for raw_line in lines:
        line = " ".join(raw_line.split()).strip(" •\t-")
        if not line:
            continue

        if not merged:
            merged.append(line)
            continue

        previous = merged[-1]
        trimmed = re.sub(r"^[\s\[\(\{\-•\"'“”‘’]+", "", line)
        starts_requirement = bool(
            re.match(r"^(shall|must|should|required|may)\b", trimmed, re.IGNORECASE)
        )
        starts_lowercase = bool(re.match(r"^[a-z]", trimmed))

        if not re.search(r"[.!?]$", previous) and (
            starts_requirement or starts_lowercase
        ):
            merged[-1] = f"{previous} {line}".strip()
        else:
            merged.append(line)

    return merged


def find_candidate_requirements(
    pages: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Find lines that may express requirements.

    These are heuristic candidates, not approved requirements.
    """
    pattern = re.compile(
        r"\b(shall\s+not|shall|must\s+not|must|required\s+to|should)\b",
        re.IGNORECASE,
    )

    candidates = []

    for page in pages:
        logical_lines = _merge_continuation_lines(page["text"].splitlines())
        for line_number, line in enumerate(logical_lines, start=1):
            if len(line) < 18 or not pattern.search(line):
                continue

            candidates.append(
                {
                    "candidate_id": f"CAND-{len(candidates) + 1:03d}",
                    "text": line,
                    "source_page": page["page"],
                    "source_line": line_number,
                    "detection_method": "rule_based",
                    "review_status": "needs_review",
                }
            )

    return candidates


def analyze_document(
    filename: str,
    data: bytes,
) -> dict[str, Any]:
    """Extract text and return preliminary requirement candidates."""
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        pages = extract_pdf(data)
    elif extension == ".docx":
        pages = extract_docx(data)
    else:
        raise ValueError("Only PDF and DOCX files are supported.")

    full_text = "\n\n".join(
        (f"[Page {page['page']}]\n" if page["page"] is not None else "") + page["text"]
        for page in pages
        if page["text"]
    )

    if not full_text.strip():
        raise ValueError("No selectable text found. The document may be scanned.")

    candidates = find_candidate_requirements(pages)

    return {
        "filename": filename,
        "page_count": sum(1 for page in pages if page["page"] is not None),
        "character_count": len(full_text),
        "text": full_text,
        "pages": pages,
        "candidate_count": len(candidates),
        "candidates": candidates,
        "review_status": "needs_review",
    }
